"""Erzeugt src/data/historical-eur.json: Jahresrenditen aus Sicht eines Anlegers in DM bzw. Euro.

Aufruf (Quelldateien zuvor herunterladen, siehe docs/DATA_SOURCES.md):

    python -I scripts/build_historical_eur.py QUELLVERZEICHNIS OUT.json

Erwartete Dateien im Quellverzeichnis:

    akkumula.html      fynup.at-Seite des DWS (ESG) Akkumula LC (DE0008474024)
    msci-world-net-eur.json  MSCI World Net Return EUR (990100, NETR), Monatsendstände
    umlauf-oeff.csv    Bundesbank BBSIS.M.I.UMR.RD.EUR.S13.B.A.A.R.A.A._Z._Z.A
    umlauf-oeff-10.csv Bundesbank BBSIS.M.I.UMR.RD.EUR.S13.B.A.R0910.R.A.A._Z._Z.A
    geld-ffm-3m.csv    Bundesbank BBIG1.M.D0.EUR.MMKT.GMFFM.M03.AVE.MA
    euribor-3m.csv     Bundesbank BBIG1.M.D0.EUR.MMKT.EURIBOR.M03.AVE.MA
    vpi-de.csv         Bundesbank BBDP1.M.DE.N.VPI.C.A00000.I20.A
    vpi-lange-reihen.xlsx  Destatis, Verbraucherpreisindex lange Reihen ab 1948

Benötigt das Python-Paket openpyxl.

Konvention: Das Kalenderjahr Y reicht vom Jahresende Y−1 bis zum Jahresende Y. Aktien nutzen
Stichtagswerte, Anleihen und Geldmarkt Monatsdurchschnitte (Dezember Y−1 bis Dezember Y), die
Inflation den Dezemberindex. Alle Renditen sind nominal und einfach (0.05 = 5 %).
"""

import csv
import hashlib
import json
import os
import re
import sys
from datetime import date

import openpyxl

FIRST_YEAR = 1963  # erster Dezember-Wert des VPI „alle privaten Haushalte“ (früheres Bundesgebiet): 1962
LAST_YEAR = 2025
MSCI_FROM = 2001  # MSCI World Net EUR liegt ab Ende Dezember 2000 vor
EURIBOR_FROM = 1999
VPI_GERMANY_FROM = 1992  # VPI Deutschland ab Januar 1991, Jahresrendite ab Dezember 1991
BOND_MATURITY = 10


def sha256(path):
    with open(path, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def read_bundesbank(path):
    """Monatsreihe aus einer Bundesbank-CSV als dict (Jahr, Monat) -> Wert; Lücken entfallen."""
    out = {}
    with open(path, encoding='utf-8-sig', newline='') as f:
        for row in csv.reader(f):
            if len(row) < 2 or not re.fullmatch(r'\d{4}-\d{2}', row[0]):
                continue
            try:
                value = float(row[1])
            except ValueError:
                continue
            y, m = row[0].split('-')
            out[(int(y), int(m))] = value
    return out


def read_akkumula(path):
    """Letzter Anteilwert je Kalenderjahr aus den Chartpunkten der fynup-Seite."""
    html = open(path, encoding='utf-8').read()
    points = re.findall(r'\\"x\\":-?\d+,\\"y\\":([0-9.]+),\\"name\\":\\"(\d\d)\.(\d\d)\.(\d{4})\\"', html)
    if not points:
        raise ValueError('Keine Akkumula-Kurspunkte gefunden')
    last = {}
    for value, day, month, year in points:
        key = (int(year), int(month), int(day))
        if int(year) not in last or key > last[int(year)][0]:
            last[int(year)] = (key, float(value))
    for year, ((_, month, day), _) in last.items():
        if year < date.today().year and (month, day) < (12, 20):
            raise ValueError(f'Akkumula: letzter Wert {year} erst am {day}.{month}.')
    return {year: value for year, (_, value) in last.items()}


def read_msci(path):
    """Stand zum letzten Handelstag im Dezember je Jahr."""
    levels = json.load(open(path, encoding='utf-8'))['indexes']['INDEX_LEVELS']
    out = {}
    for item in levels:
        d = str(item['calc_date'])
        if d[4:6] == '12':
            out[int(d[:4])] = float(item['level_eod'])
    return out


def read_vpi_west(path):
    """Destatis, Blatt „Monate“: Preisindex für die Lebenshaltung aller privaten Haushalte,
    früheres Bundesgebiet (1995 = 100), Dezemberwerte."""
    sheet = openpyxl.load_workbook(path, data_only=True)['Monate']
    header = [c for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    if header[4] != 'Früheres Bundesgebiet':
        raise ValueError('Destatis-Tabelle: Spalte „Früheres Bundesgebiet“ nicht an erwarteter Stelle')
    out, year = {}, None
    for row in sheet.iter_rows(min_row=6, values_only=True):
        if isinstance(row[0], int):
            year = row[0]
        if row[1] == 'Dezember' and isinstance(row[4], (int, float)):
            out[year] = float(row[4])
    return out


def bond_month(y0, y1, maturity):
    """Monatsrendite einer Anleihe mit konstanter Laufzeit: zum Rendite-Niveau y0 zu pari gekauft
    (Kupon y0), einen Monat später mit Restlaufzeit maturity − 1/12 zur Rendite y1 bewertet."""
    t = maturity - 1 / 12
    discount = (1 + y1) ** -t
    price = y0 / y1 * (1 - discount) + discount if y1 != 0 else 1 + y0 * t
    return y0 / 12 + price - 1


def main(src, out_path):
    p = lambda name: os.path.join(src, name)  # noqa: E731
    akkumula = read_akkumula(p('akkumula.html'))
    msci = read_msci(p('msci-world-net-eur.json'))
    yields_all = read_bundesbank(p('umlauf-oeff.csv'))
    yields_10 = read_bundesbank(p('umlauf-oeff-10.csv'))
    money_ffm = read_bundesbank(p('geld-ffm-3m.csv'))
    euribor = read_bundesbank(p('euribor-3m.csv'))
    vpi_de = read_bundesbank(p('vpi-de.csv'))
    vpi_west = read_vpi_west(p('vpi-lange-reihen.xlsx'))

    def bond_step(prev, cur):
        # Beide Monate aus derselben Reihe, sonst erzeugt der Reihenwechsel einen Kurssprung
        series = yields_10 if prev in yields_10 and cur in yields_10 else yields_all
        return bond_month(series[prev] / 100, series[cur] / 100, BOND_MATURITY)

    rows = []
    for year in range(FIRST_YEAR, LAST_YEAR + 1):
        stock = akkumula if year < MSCI_FROM else msci
        equity = stock[year] / stock[year - 1] - 1

        months = [(year - 1, 12)] + [(year, m) for m in range(1, 13)]
        bond = 1.0
        for prev, cur in zip(months, months[1:]):
            bond *= 1 + bond_step(prev, cur)

        money = euribor if year >= EURIBOR_FROM else money_ffm
        cash = 1.0
        for m in range(1, 13):
            cash *= 1 + money[(year, m)] / 1200

        if year >= VPI_GERMANY_FROM:
            inflation = vpi_de[(year, 12)] / vpi_de[(year - 1, 12)] - 1
        else:
            inflation = vpi_west[year] / vpi_west[year - 1] - 1

        rows.append({
            'year': year,
            'equity': round(equity, 8),
            'bonds': round(bond - 1, 8),
            'cash': round(cash - 1, 8),
            'inflation': round(inflation, 8),
        })

    # Abstände der verketteten Reihen in den Überlappungszeiträumen (Prozentpunkte)
    def mean_spread(a, b, first, last):
        keys = [k for k in a if k in b and first <= k <= last]
        return round(sum(a[k] - b[k] for k in keys) / len(keys), 3)

    splices = {
        'bondYield10yMinusAllMaturities1973to2025': mean_spread(
            yields_10, yields_all, (1973, 4), (LAST_YEAR, 12)),
        'euriborMinusFrankfurt3m1999to2012': mean_spread(euribor, money_ffm, (1999, 1), (2012, 5)),
    }

    bundesbank = 'https://api.statistiken.bundesbank.de/rest/data/'
    data = {
        'meta': {
            'description': (
                'Jährliche nominale Renditen aus Sicht eines Anlegers in DM bzw. Euro: Aktien bis '
                '2000 DWS Akkumula (Anteilwert nach Fondskosten), ab 2001 MSCI World Net Return EUR; '
                'rechnerische 10-jährige deutsche Staatsanleihe (Total Return) aus der '
                'Umlaufsrendite; Dreimonatsgeld (Frankfurt, ab 1999 Euribor); VPI-Inflation '
                '(bis 1991 früheres Bundesgebiet).'
            ),
            'convention': 'Jahr Y = Jahresende Y−1 bis Jahresende Y',
            'generated': date.today().isoformat(),
            'firstYear': rows[0]['year'],
            'lastYear': rows[-1]['year'],
            'equitySplice': {'msciFromYear': MSCI_FROM},
            'splices': splices,
            'sources': [
                {
                    'name': 'DWS (ESG) Akkumula LC, DE0008474024, Anteilwerte via fynup.at (Aktien bis 2000)',
                    'url': 'https://www.fynup.at/f/de0008474024',
                    'sha256': sha256(p('akkumula.html')),
                },
                {
                    'name': 'MSCI World Index (990100), Net Return, EUR, Monatsendstände (Aktien ab 2001)',
                    'url': 'https://app2.msci.com/products/service/index/indexmaster/getLevelDataForGraph'
                    '?currency_symbol=EUR&index_variant=NETR&start_date=19970101&end_date=20251231'
                    '&data_frequency=END_OF_MONTH&index_codes=990100',
                    'sha256': sha256(p('msci-world-net-eur.json')),
                },
                {
                    'name': 'Deutsche Bundesbank, Umlaufsrendite öffentliche Hand (Anleihen bis 03/1973 und 03/1976–01/1977)',
                    'url': bundesbank + 'BBSIS/M.I.UMR.RD.EUR.S13.B.A.A.R.A.A._Z._Z.A',
                    'sha256': sha256(p('umlauf-oeff.csv')),
                },
                {
                    'name': 'Deutsche Bundesbank, Umlaufsrendite öffentliche Hand, Restlaufzeit 9–10 Jahre (Anleihen ab 04/1973, Lücke 03–12/1976)',
                    'url': bundesbank + 'BBSIS/M.I.UMR.RD.EUR.S13.B.A.R0910.R.A.A._Z._Z.A',
                    'sha256': sha256(p('umlauf-oeff-10.csv')),
                },
                {
                    'name': 'Deutsche Bundesbank, Dreimonatsgeld Frankfurter Banken, Monatsdurchschnitt (Cash bis 1998)',
                    'url': bundesbank + 'BBIG1/M.D0.EUR.MMKT.GMFFM.M03.AVE.MA',
                    'sha256': sha256(p('geld-ffm-3m.csv')),
                },
                {
                    'name': 'Deutsche Bundesbank, Dreimonats-Euribor, Monatsdurchschnitt (Cash ab 1999)',
                    'url': bundesbank + 'BBIG1/M.D0.EUR.MMKT.EURIBOR.M03.AVE.MA',
                    'sha256': sha256(p('euribor-3m.csv')),
                },
                {
                    'name': 'Statistisches Bundesamt, Verbraucherpreisindex Deutschland via Bundesbank (Inflation ab 1992)',
                    'url': bundesbank + 'BBDP1/M.DE.N.VPI.C.A00000.I20.A',
                    'sha256': sha256(p('vpi-de.csv')),
                },
                {
                    'name': 'Statistisches Bundesamt, Verbraucherpreisindex lange Reihen ab 1948, früheres Bundesgebiet (Inflation bis 1991)',
                    'url': 'https://www.destatis.de/DE/Themen/Wirtschaft/Preise/Verbraucherpreisindex/Publikationen/Downloads-Verbraucherpreise/verbraucherpreisindex-lange-reihen-pdf-5611103.html',
                    'sha256': sha256(p('vpi-lange-reihen.xlsx')),
                },
            ],
        },
        'rows': rows,
    }
    # Metadaten eingerückt, Jahreszeilen je eine Zeile
    meta = json.dumps(data['meta'], ensure_ascii=False, indent=2).replace('\n', '\n  ')
    lines = ',\n'.join('    ' + json.dumps(row, ensure_ascii=False) for row in rows)
    with open(out_path, 'w', encoding='utf-8') as f:
        f.write('{\n  "meta": ' + meta + ',\n  "rows": [\n' + lines + '\n  ]\n}\n')
    print(f'{len(rows)} Jahre {rows[0]["year"]}–{rows[-1]["year"]}, Spreads {splices}')


if __name__ == '__main__':
    main(*sys.argv[1:3])
